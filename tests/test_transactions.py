import io
from datetime import date

from openpyxl import Workbook, load_workbook

from .conftest import API, category_id, h, make_user


def add_tx(client, user, **kw):
    body = {"kind": "expense", "amount": "10.00", "date": "2026-03-01", "description": "x"} | kw
    r = client.post(f"{API}/transactions", json=body, headers=h(user))
    assert r.status_code == 201, r.text
    return r.json()


def test_create_validates_amount_and_kind(client, user):
    for bad in ({"amount": "-5"}, {"amount": "0"}, {"amount": "1.234"}, {"kind": "transfer"}, {"date": "not-a-date"}):
        body = {"kind": "expense", "amount": "10", "date": "2026-01-01"} | bad
        assert client.post(f"{API}/transactions", json=body, headers=h(user)).status_code == 422


def test_category_kind_must_match(client, user):
    salary = category_id(client, user, "Salary", "income")
    r = client.post(f"{API}/transactions", json={"kind": "expense", "amount": "5", "date": "2026-01-01",
                                                 "category_id": salary}, headers=h(user))
    assert r.status_code == 400 and r.json()["error"]["code"] == "category_kind_mismatch"


def test_search_filter_sort_paginate(client, user):
    groceries = category_id(client, user, "Groceries")
    add_tx(client, user, description="Coffee beans", amount="12.50", date="2026-01-05", category_id=groceries)
    add_tx(client, user, description="Weekly groceries", amount="80.00", date="2026-01-10", category_id=groceries)
    add_tx(client, user, description="Salary", amount="3000", kind="income", date="2026-01-01")
    add_tx(client, user, description="100% juice_box", amount="3", date="2026-02-01")

    r = client.get(f"{API}/transactions", params={"q": "coffee"}, headers=h(user)).json()
    assert r["total"] == 1 and r["items"][0]["description"] == "Coffee beans"
    # LIKE wildcards in the search term are matched literally.
    assert client.get(f"{API}/transactions", params={"q": "%"}, headers=h(user)).json()["total"] == 1
    assert client.get(f"{API}/transactions", params={"kind": "income"}, headers=h(user)).json()["total"] == 1
    r = client.get(f"{API}/transactions", params={"category_id": groceries, "sort": "amount", "order": "asc"},
                   headers=h(user)).json()
    assert [i["amount"] for i in r["items"]] == ["12.50", "80.00"]
    r = client.get(f"{API}/transactions", params={"date_from": "2026-01-06", "date_to": "2026-01-31"}, headers=h(user)).json()
    assert r["total"] == 1
    r = client.get(f"{API}/transactions", params={"min_amount": 50}, headers=h(user)).json()
    assert r["total"] == 2
    page = client.get(f"{API}/transactions", params={"size": 2, "page": 2}, headers=h(user)).json()
    assert page["pages"] == 2 and len(page["items"]) == 2
    assert client.get(f"{API}/transactions", params={"sort": "password_hash"}, headers=h(user)).status_code == 422


def test_update_and_delete(client, user):
    t = add_tx(client, user)
    r = client.patch(f"{API}/transactions/{t['id']}", json={"amount": "25.00", "description": "edited"}, headers=h(user))
    assert r.status_code == 200 and r.json()["amount"] == "25.00"
    assert client.delete(f"{API}/transactions/{t['id']}", headers=h(user)).status_code == 204
    assert client.get(f"{API}/transactions/{t['id']}", headers=h(user)).status_code == 404


def test_users_cannot_access_each_others_transactions(client, user):
    other = make_user(client)
    t = add_tx(client, user)
    assert client.get(f"{API}/transactions/{t['id']}", headers=h(other)).status_code == 404
    assert client.patch(f"{API}/transactions/{t['id']}", json={"amount": "1"}, headers=h(other)).status_code == 404
    assert client.delete(f"{API}/transactions/{t['id']}", headers=h(other)).status_code == 404
    assert client.get(f"{API}/transactions", headers=h(other)).json()["total"] == 0
    # Using someone else's category is also rejected.
    their_cat = category_id(client, user, "Dining")
    r = client.post(f"{API}/transactions", json={"kind": "expense", "amount": "5", "date": "2026-01-01",
                                                 "category_id": their_cat}, headers=h(other))
    assert r.status_code == 404


def test_csv_import_and_export(client, user):
    csv_data = ("date,kind,amount,category,description\n"
                "2026-01-02,expense,42.10,Groceries,Market\n"
                "2026-01-03,,-15.00,Coffee Shops,Latte\n"   # sign infers kind; unknown category is created
                "2026-01-04,income,1000,Salary,Pay\n"
                "bad-date,expense,10,,Broken row\n"
                "2026-01-05,expense,abc,,Bad amount\n")
    r = client.post(f"{API}/transactions/import", files={"file": ("tx.csv", csv_data, "text/csv")}, headers=h(user))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["imported"] == 3 and body["skipped"] == 2 and {e["row"] for e in body["errors"]} == {5, 6}
    assert any(c["name"] == "Coffee Shops" for c in client.get(f"{API}/categories", headers=h(user)).json())

    add_tx(client, user, description="=HYPERLINK(\"http://evil\")")
    r = client.get(f"{API}/transactions/export", params={"format": "csv"}, headers=h(user))
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    text = r.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text  # formula injection neutralised
    assert text.count("\n") == 5


def test_xlsx_import_and_export(client, user):
    wb = Workbook()
    ws = wb.active
    ws.append(["date", "kind", "amount", "description"])
    ws.append([date(2026, 2, 1), "expense", 19.99, "Book"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post(f"{API}/transactions/import", headers=h(user),
                    files={"file": ("tx.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.json()["imported"] == 1
    r = client.get(f"{API}/transactions/export", params={"format": "xlsx"}, headers=h(user))
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert rows[0][0] == "date" and rows[1][2] == 19.99


def test_upload_rejects_bad_files(client, user):
    r = client.post(f"{API}/transactions/import", files={"file": ("x.exe", b"MZ", "application/octet-stream")}, headers=h(user))
    assert r.status_code == 415
    r = client.post(f"{API}/transactions/import", files={"file": ("x.xlsx", b"not a zip", "application/octet-stream")}, headers=h(user))
    assert r.status_code == 400
    big = b"date,kind,amount\n" + b"2026-01-01,expense,1\n" * 300_000
    r = client.post(f"{API}/transactions/import", files={"file": ("big.csv", big, "text/csv")}, headers=h(user))
    assert r.status_code == 413


def test_error_envelope_on_404(client, user):
    r = client.get(f"{API}/transactions/999999", headers=h(user))
    assert r.status_code == 404
    err = r.json()["error"]
    assert err["code"] == "not_found" and err["request_id"]

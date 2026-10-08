"""CSV / Excel import and export.

Upload hardening: size cap enforced while streaming, extension + magic-byte check, row cap,
every row validated through the same Pydantic schema as the JSON API, and spreadsheet
formula injection neutralised on export."""
import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.errors import AppError, PayloadTooLarge
from app.models import Category, Transaction
from app.schemas.finance import TransactionIn

EXPORT_COLUMNS = ["date", "kind", "amount", "category", "description", "notes"]
MAX_NEW_CATEGORIES = 50
ZIP_MAGIC = b"PK\x03\x04"


async def read_upload(file: UploadFile) -> tuple[bytes, str]:
    name = (file.filename or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    if ext not in ("csv", "xlsx"):
        raise AppError("Only .csv and .xlsx files are supported", code="unsupported_file_type", status_code=415)
    chunks, size = [], 0
    while chunk := await file.read(64 * 1024):
        size += len(chunk)
        if size > settings.max_upload_bytes:
            raise PayloadTooLarge(f"File exceeds {settings.max_upload_bytes // (1024 * 1024)} MB limit")
        chunks.append(chunk)
    data = b"".join(chunks)
    if not data:
        raise AppError("File is empty", code="empty_file")
    if ext == "xlsx" and not data.startswith(ZIP_MAGIC):
        raise AppError("File content is not a valid .xlsx workbook", code="invalid_file")
    if ext == "csv" and (b"\x00" in data[:4096] or data.startswith(ZIP_MAGIC)):
        raise AppError("File content is not valid CSV text", code="invalid_file")
    return data, ext


def _rows_from_csv(data: bytes):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise AppError("CSV must be UTF-8 encoded", code="invalid_file") from None
    reader = csv.DictReader(io.StringIO(text))
    for row in reader:
        yield {(k or "").strip().lower(): (v or "").strip() for k, v in row.items() if k}


def _rows_from_xlsx(data: bytes):
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise AppError("Could not read the Excel workbook", code="invalid_file") from None
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows, [])]
    for values in rows:
        if values and any(v not in (None, "") for v in values):
            yield {h: v for h, v in zip(header, values, strict=False) if h}
    wb.close()


def _parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Unrecognised date '{s}' (use YYYY-MM-DD)")


def _parse_row(row: dict) -> dict:
    raw_amount = str(row.get("amount", "")).replace(",", "").replace("$", "").strip()
    try:
        amount = Decimal(raw_amount)
    except InvalidOperation:
        raise ValueError(f"Invalid amount '{row.get('amount')}'") from None
    kind = str(row.get("kind") or row.get("type") or "").strip().lower()
    if not kind:
        kind = "expense" if amount < 0 else "income"
    return {"kind": kind, "amount": abs(amount), "date": _parse_date(row.get("date")),
            "description": str(row.get("description") or "")[:255], "notes": str(row.get("notes") or "")[:2000],
            "category": str(row.get("category") or "").strip()[:60]}


def import_transactions(db: Session, user_id: int, data: bytes, ext: str) -> dict:
    cats = {(c.name.lower(), c.kind): c for c in db.scalars(select(Category).where(Category.user_id == user_id))}
    new_cats, imported, errors, skipped = 0, 0, [], 0
    rows = _rows_from_csv(data) if ext == "csv" else _rows_from_xlsx(data)
    for line, row in enumerate(rows, start=2):  # line 1 is the header
        if line - 1 > settings.max_import_rows:
            raise AppError(f"Too many rows (max {settings.max_import_rows})", code="too_many_rows")
        try:
            parsed = _parse_row(row)
            cat_name = parsed.pop("category")
            payload = TransactionIn(**parsed)
        except (ValueError, ValidationError) as exc:
            skipped += 1
            if len(errors) < 50:
                msg = exc.errors()[0]["msg"] if isinstance(exc, ValidationError) else str(exc)
                errors.append({"row": line, "error": msg})
            continue
        category = None
        if cat_name:
            category = cats.get((cat_name.lower(), payload.kind))
            if category is None and new_cats < MAX_NEW_CATEGORIES:
                category = Category(user_id=user_id, name=cat_name, kind=payload.kind)
                db.add(category)
                db.flush()
                cats[(cat_name.lower(), payload.kind)] = category
                new_cats += 1
        db.add(Transaction(user_id=user_id, category_id=category.id if category else None,
                           **payload.model_dump(exclude={"category_id"})))
        imported += 1
    return {"imported": imported, "skipped": skipped, "errors": errors}


def _safe_cell(v):
    """Prevent CSV/formula injection when the export is opened in a spreadsheet."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _export_rows(txs):
    for t in txs:
        yield [t.date.isoformat(), t.kind, str(t.amount), t.category.name if t.category else "", t.description, t.notes]


def export_csv(txs) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(EXPORT_COLUMNS)
    for r in _export_rows(txs):
        w.writerow([_safe_cell(v) for v in r])
    return buf.getvalue().encode("utf-8-sig")


def export_xlsx(txs) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Transactions")
    ws.append(EXPORT_COLUMNS)
    for d, kind, amount, cat, desc, notes in _export_rows(txs):
        ws.append([date.fromisoformat(d), kind, float(amount), _safe_cell(cat), _safe_cell(desc), _safe_cell(notes)])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template_csv() -> bytes:
    return (b"date,kind,amount,category,description,notes\n"
            b"2026-01-05,expense,42.50,Groceries,Weekly groceries,\n"
            b"2026-01-01,income,3200.00,Salary,Monthly salary,\n")
